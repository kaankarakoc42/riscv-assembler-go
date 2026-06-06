from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
IN_DOCX = ROOT / "build" / "reports" / "BIL302_PROJE3_RV32I_LOADER_RAPORU.docx"
OUT_PDF = ROOT / "build" / "reports" / "BIL302_PROJE3_RV32I_LOADER_RAPORU.pdf"


def iter_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            for p in doc.paragraphs:
                if p._p is child:
                    yield ("p", p)
                    break
        elif child.tag == qn("w:tbl"):
            for t in doc.tables:
                if t._tbl is child:
                    yield ("t", t)
                    break


def style_for(paragraph, styles):
    name = paragraph.style.name if paragraph.style is not None else ""
    text = paragraph.text.strip()
    if not text:
        return None
    if name == "Title":
        return styles["title"]
    if name == "Heading 1":
        return styles["h1"]
    if name == "Heading 2":
        return styles["h2"]
    if name == "Heading 3":
        return styles["h3"]
    if name == "Code":
        return styles["code"]
    if name == "List Bullet":
        return styles["bullet"]
    if name == "List Number":
        return styles["bullet"]
    return styles["body"]


def make_pdf() -> None:
    pdfmetrics.registerFont(TTFont("CourierNew", r"C:\Windows\Fonts\cour.ttf"))
    pdfmetrics.registerFont(TTFont("CourierNew-Bold", r"C:\Windows\Fonts\courbd.ttf"))
    pdfmetrics.registerFont(TTFont("CourierNew-Italic", r"C:\Windows\Fonts\couri.ttf"))

    styles = {
        "body": ParagraphStyle("body", fontName="CourierNew", fontSize=10, leading=12, spaceAfter=6, alignment=TA_LEFT),
        "title": ParagraphStyle("title", fontName="CourierNew-Bold", fontSize=14, leading=17, spaceAfter=10, alignment=TA_CENTER, textColor=colors.HexColor("#0B2545")),
        "h1": ParagraphStyle("h1", fontName="CourierNew-Bold", fontSize=12, leading=15, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#1F4D78")),
        "h2": ParagraphStyle("h2", fontName="CourierNew-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor("#2E74B5")),
        "h3": ParagraphStyle("h3", fontName="CourierNew-Bold", fontSize=10, leading=12, spaceBefore=6, spaceAfter=3, textColor=colors.HexColor("#1F4D78")),
        "code": ParagraphStyle("code", fontName="CourierNew", fontSize=8.5, leading=10, leftIndent=12, spaceAfter=6),
        "bullet": ParagraphStyle("bullet", fontName="CourierNew", fontSize=10, leading=12, leftIndent=18, firstLineIndent=-10, spaceAfter=3),
        "cell": ParagraphStyle("cell", fontName="CourierNew", fontSize=7.6, leading=9.2, spaceAfter=0),
        "cell_bold": ParagraphStyle("cell_bold", fontName="CourierNew-Bold", fontSize=7.6, leading=9.2, spaceAfter=0),
    }

    docx = Document(IN_DOCX)
    story = []
    numbered = 1
    usable_width = 6.5 * inch

    for kind, block in iter_blocks(docx):
        if kind == "p":
            st = style_for(block, styles)
            if st is None:
                continue
            text = block.text
            name = block.style.name if block.style is not None else ""
            if name == "Code":
                story.append(Preformatted(text.rstrip(), st, maxLineLength=88))
            elif name == "List Bullet":
                story.append(Paragraph("• " + escape(text), st))
            elif name == "List Number":
                story.append(Paragraph(f"{numbered}. " + escape(text), st))
                numbered += 1
            else:
                if name not in {"List Number"}:
                    numbered = 1
                if text.startswith("Şekil 1."):
                    img = ROOT / "docs" / "images" / "mermaid-jpeg" / "diagram-02-line-50.jpeg"
                    if img.exists():
                        story.append(Image(str(img), width=5.7 * inch, height=2.75 * inch, hAlign="CENTER"))
                        story.append(Spacer(1, 4))
                story.append(Paragraph(escape(text).replace("\n", "<br/>"), st))
            continue

        data = []
        col_count = len(block.rows[0].cells) if block.rows else 0
        for r_idx, row in enumerate(block.rows):
            row_data = []
            for cell in row.cells:
                txt = escape(cell.text.strip()).replace("\n", "<br/>")
                row_data.append(Paragraph(txt, styles["cell_bold" if r_idx == 0 else "cell"]))
            data.append(row_data)
        if not data or col_count == 0:
            continue
        col_width = usable_width / col_count
        tbl = Table(data, colWidths=[col_width] * col_count, repeatRows=1, hAlign="CENTER")
        tbl.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#999999")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 8))

    def page(canvas, _doc):
        canvas.saveState()
        canvas.setFont("CourierNew", 8)
        canvas.drawString(inch, 0.55 * inch, "BIL302 Proje 3 - RV32I FPGA Loader")
        canvas.drawRightString(7.5 * inch, 0.55 * inch, f"Sayfa {_doc.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=letter,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="BIL302 Proje 3 RV32I Loader Raporu",
        author="BIL302 Proje Grubu",
    )
    pdf.build(story, onFirstPage=page, onLaterPages=page)
    print(OUT_PDF)


if __name__ == "__main__":
    make_pdf()
