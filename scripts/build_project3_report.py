from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from textwrap import wrap
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "build" / "reports"
ASSET = ROOT / "build" / "report-assets"
DOT_DIR = ASSET / "dot"
SVG_DIR = ASSET / "svg"
PNG_DIR = ASSET / "png"

DOCX = OUT / "BIL302_PROJE3_A.XX_B.YY_C.ZZ_D.WW_170526.docx"
PDF = OUT / "BIL302_PROJE3_A.XX_B.YY_C.ZZ_D.WW_170526.PDF"

FONT = "Courier New"
BLUE = "1F4D78"
INK = "0B2545"
FILL = "E8EEF5"


@dataclass
class Figure:
    slug: str
    title: str
    dot: str
    boxes: list[tuple[str, int, int, int, int]]
    edges: list[tuple[int, int]]
    size: tuple[int, int] = (1240, 620)


FIGURES = [
    Figure(
        "toolchain",
        "Şekil 1. Assembler, linker, host uygulaması ve FPGA loader uçtan uca veri akışı.",
        """
digraph toolchain {
  rankdir=LR;
  source[label="RV32I .s kaynak"];
  asm[label="rvasm / RVOB1"];
  link[label="rvld / .bin .hex .mem"];
  host[label="host_loader.py"];
  uart[label="UART paketleri + checksum"];
  fpga[label="PicoRV32 loader + BRAM"];
  source -> asm -> link -> host -> uart -> fpga;
}
""".strip(),
        [
            ("RV32I .s\nkaynak", 40, 230, 170, 90),
            ("rvasm\nRVOB1", 240, 230, 160, 90),
            ("rvld\n.bin/.hex/.mem", 460, 230, 190, 90),
            ("host_loader.py", 720, 230, 190, 90),
            ("UART paketleri\nchecksum", 970, 230, 190, 90),
            ("PicoRV32 loader\nBRAM", 970, 430, 190, 90),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
    ),
    Figure(
        "compiler_pipeline",
        "Şekil 2. rvasm compiler/assembler tasarım boru hattı.",
        """
digraph compiler_pipeline {
  rankdir=LR;
  source -> lexer -> parser -> pass1 -> pass2 -> rvob1;
  source[label="Assembly source"];
  lexer[label="Lexer"];
  parser[label="Parser / AST"];
  pass1[label="Pass 1\\nsection + symbols"];
  pass2[label="Pass 2\\nencoding + reloc"];
  rvob1[label="RVOB1 object"];
}
""".strip(),
        [
            ("Assembly\nsource", 35, 235, 150, 90),
            ("Lexer\ntoken stream", 230, 235, 170, 90),
            ("Parser\nAST", 445, 235, 150, 90),
            ("Pass 1\nsection + symbols", 640, 235, 220, 90),
            ("Pass 2\nencoding + reloc", 900, 235, 220, 90),
            ("RVOB1\nobject", 520, 430, 160, 90),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
    ),
    Figure(
        "compiler_data_model",
        "Şekil 3. Compiler/assembler içinde kullanılan ana veri yapıları.",
        """
digraph compiler_data_model {
  rankdir=LR;
  ast -> sections;
  ast -> symbols;
  ast -> relocations;
  sections -> rvob1;
  symbols -> rvob1;
  relocations -> rvob1;
}
""".strip(),
        [
            ("AST\nInstruction/Directive/Label", 60, 250, 260, 100),
            (".text/.data\nsection buffer", 460, 100, 260, 90),
            ("Symbol table\nlocal/global/extern", 460, 250, 260, 90),
            ("Relocation list\nbranch/jal/hi-lo", 460, 400, 260, 90),
            ("RVOB1 writer\nheader + tables", 860, 250, 260, 100),
        ],
        [(0, 1), (0, 2), (0, 3), (1, 4), (2, 4), (3, 4)],
    ),
    Figure(
        "linker_pipeline",
        "Şekil 4. rvld linker mimarisi ve final imaj üretim akışı.",
        """
digraph linker_pipeline {
  rankdir=LR;
  inputs -> script -> merge -> symtab -> reloc -> emit;
  inputs[label="RVOB1 inputs"];
  script[label="link.toml memory"];
  merge[label="section merge"];
  symtab[label="global symbol table"];
  reloc[label="relocation pass"];
  emit[label="bin/hex/mem emit"];
}
""".strip(),
        [
            ("RVOB1\ninputlar", 40, 230, 160, 90),
            ("link.toml\nmemory layout", 240, 230, 190, 90),
            ("Section merge\n.text + .data", 480, 230, 210, 90),
            ("Global symbol\ntable", 740, 230, 190, 90),
            ("Relocation\npatch", 970, 140, 170, 90),
            ("Image emit\n.bin/.hex/.mem", 970, 350, 200, 90),
        ],
        [(0, 2), (1, 2), (2, 3), (3, 4), (4, 5)],
    ),
    Figure(
        "relocation_flow",
        "Şekil 5. Linker relocation patch karar akışı.",
        """
digraph relocation_flow {
  rankdir=LR;
  read -> calc -> choose -> patch -> range -> write;
  range -> error [label="taşma"];
}
""".strip(),
        [
            ("Reloc kaydı\noku", 35, 235, 160, 90),
            ("S + A - P\nhesapla", 230, 235, 170, 90),
            ("Türe göre\nimmediate seç", 440, 235, 200, 90),
            ("Bit alanlarını\npatch et", 680, 235, 190, 90),
            ("Range check", 910, 235, 160, 90),
            ("Final word\nimage'a yaz", 560, 430, 190, 90),
            ("Hard error\nimage yok", 910, 430, 180, 90),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (4, 6)],
    ),
    Figure(
        "loader_fsm",
        "Şekil 6. FPGA tarafındaki software-loader durum makinesi.",
        """
digraph loader_fsm {
  rankdir=LR;
  READY -> MAGIC -> HEADER -> DATA -> CHECKSUM;
  CHECKSUM -> ACK [label="doğru"];
  CHECKSUM -> ERR [label="hatalı"];
  ACK -> READY;
  ACK -> START [label="cmd=START"];
}
""".strip(),
        [
            ("READY\n'R' gönder", 40, 230, 160, 80),
            ("MAGIC\n0x55 bekle", 230, 230, 160, 80),
            ("HEADER\ncmd+addr+words", 420, 230, 190, 80),
            ("DATA\nword oku", 650, 230, 160, 80),
            ("CHECKSUM\nlow8 karşılaştır", 840, 230, 210, 80),
            ("ACK\n'K'/'S'", 1060, 120, 140, 80),
            ("ERR\n'E'", 1060, 340, 140, 80),
            ("START\njr entry", 650, 430, 160, 80),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5), (4, 6), (5, 0), (5, 7)],
    ),
    Figure(
        "memory_map",
        "Şekil 7. Loader, kullanıcı program alanı ve UART MMIO bellek haritası.",
        """
digraph memory_map {
  rankdir=TB;
  boot[label="0x00000000-0x000003FF\\nBoot loader"];
  user[label="0x00000400-0x00001FFF\\nKullanıcı programı"];
  rx[label="0x80000008 / 0x8000000C\\nUART RX"];
  tx[label="0x80000010 / 0x80000014\\nUART TX"];
  boot -> user -> rx -> tx;
}
""".strip(),
        [
            ("0x00000000-0x000003FF\nBoot loader (444 B)", 90, 70, 430, 90),
            ("0x00000400-0x00001FFF\nKullanıcı program alanı", 90, 190, 430, 120),
            ("0x80000008 / 0x8000000C\nUART RX_DATA / RX_READY", 680, 120, 430, 90),
            ("0x80000010 / 0x80000014\nUART TX_DATA / TX_BUSY", 680, 260, 430, 90),
        ],
        [(0, 1), (1, 2), (2, 3)],
    ),
    Figure(
        "protocol",
        "Şekil 8. Host-loader UART paket formatı.",
        """
digraph protocol {
  rankdir=LR;
  magic -> cmd -> addr -> words -> data -> checksum;
}
""".strip(),
        [
            ("Magic\n0x55", 45, 240, 140, 90),
            ("CMD\nWRITE/START", 225, 240, 170, 90),
            ("ADDR\nlittle-endian", 440, 240, 190, 90),
            ("WORDS\nadet", 675, 240, 140, 90),
            ("DATA\n32-bit word dizisi", 855, 240, 200, 90),
            ("SUM\nlow8", 1090, 240, 120, 90),
        ],
        [(0, 1), (1, 2), (2, 3), (3, 4), (4, 5)],
    ),
    Figure(
        "test_coverage",
        "Şekil 9. Deney senaryolarının kapsadığı sistem davranışları.",
        """
digraph test_coverage {
  rankdir=LR;
  blink -> io;
  counter -> calls;
  uarthello -> data;
  checksum -> error;
}
""".strip(),
        [
            ("Blink / Knight Rider\nDöngü + LED MMIO", 70, 90, 280, 90),
            ("Counter\nAlt program + sayaç", 70, 230, 280, 90),
            ("UART Hello\n.data + string", 70, 370, 280, 90),
            ("Checksum negatif\nHata yakalama", 70, 510, 280, 90),
            ("Gözlenebilir I/O", 760, 90, 280, 90),
            ("call/extern/link", 760, 230, 280, 90),
            ("la pseudo + bellek", 760, 370, 280, 90),
            ("ACK/ERR kararlılığı", 760, 510, 280, 90),
        ],
        [(0, 4), (1, 5), (2, 6), (3, 7)],
        (1240, 700),
    ),
]


def font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\courbd.ttf" if bold else r"C:\Windows\Fonts\cour.ttf",
        r"C:\Windows\Fonts\consolab.ttf" if bold else r"C:\Windows\Fonts\consola.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str) -> None:
    draw.line([start, end], fill=fill, width=4)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        pts = [(x2, y2), (x2 - 14 if x2 > x1 else x2 + 14, y2 - 8), (x2 - 14 if x2 > x1 else x2 + 14, y2 + 8)]
    else:
        pts = [(x2, y2), (x2 - 8, y2 - 14 if y2 > y1 else y2 + 14), (x2 + 8, y2 - 14 if y2 > y1 else y2 + 14)]
    draw.polygon(pts, fill=fill)


def write_figures() -> None:
    DOT_DIR.mkdir(parents=True, exist_ok=True)
    SVG_DIR.mkdir(parents=True, exist_ok=True)
    PNG_DIR.mkdir(parents=True, exist_ok=True)
    title_font = font(24, True)
    label_font = font(21)
    for fig in FIGURES:
        (DOT_DIR / f"{fig.slug}.dot").write_text(fig.dot + "\n", encoding="utf-8")
        w, h = fig.size
        svg_parts = [
            f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}">',
            '<rect width="100%" height="100%" fill="#ffffff"/>',
            f'<text x="32" y="38" font-family="Courier New, monospace" font-size="24" font-weight="700" fill="#{INK}">{escape(fig.title.split(". ", 1)[-1])}</text>',
            '<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="#1F4D78"/></marker></defs>',
        ]
        for a, b in fig.edges:
            ax, ay, aw, ah = fig.boxes[a][1:]
            bx, by, bw, bh = fig.boxes[b][1:]
            start = (ax + aw, ay + ah // 2)
            end = (bx, by + bh // 2)
            if bx < ax:
                start = (ax, ay + ah // 2)
                end = (bx + bw, by + bh // 2)
            svg_parts.append(f'<line x1="{start[0]}" y1="{start[1]}" x2="{end[0]}" y2="{end[1]}" stroke="#1F4D78" stroke-width="4" marker-end="url(#arrow)"/>')
        for text, x, y, bw, bh in fig.boxes:
            svg_parts.append(f'<rect x="{x}" y="{y}" width="{bw}" height="{bh}" rx="12" fill="#F7FAFC" stroke="#1F4D78" stroke-width="3"/>')
            lines = text.split("\n")
            for i, line in enumerate(lines):
                yy = y + bh / 2 - (len(lines) - 1) * 13 + i * 26 + 8
                svg_parts.append(f'<text x="{x + bw/2}" y="{yy:.0f}" text-anchor="middle" font-family="Courier New, monospace" font-size="21" fill="#0B2545">{escape(line)}</text>')
        svg_parts.append("</svg>")
        (SVG_DIR / f"{fig.slug}.svg").write_text("\n".join(svg_parts), encoding="utf-8")

        img = Image.new("RGB", (w, h), "white")
        draw = ImageDraw.Draw(img)
        draw.text((32, 16), fig.title.split(". ", 1)[-1], fill=f"#{INK}", font=title_font)
        for a, b in fig.edges:
            ax, ay, aw, ah = fig.boxes[a][1:]
            bx, by, bw, bh = fig.boxes[b][1:]
            start = (ax + aw, ay + ah // 2)
            end = (bx, by + bh // 2)
            if bx < ax:
                start = (ax, ay + ah // 2)
                end = (bx + bw, by + bh // 2)
            arrow(draw, start, end, f"#{BLUE}")
        for text, x, y, bw, bh in fig.boxes:
            draw.rounded_rectangle((x, y, x + bw, y + bh), radius=12, fill="#F7FAFC", outline=f"#{BLUE}", width=3)
            lines = text.split("\n")
            for i, line in enumerate(lines):
                bbox = draw.textbbox((0, 0), line, font=label_font)
                tx = x + (bw - (bbox[2] - bbox[0])) / 2
                ty = y + bh / 2 - len(lines) * 13 + i * 26
                draw.text((tx, ty), line, fill=f"#{INK}", font=label_font)
        img.save(PNG_DIR / f"{fig.slug}.png", quality=95)


def set_run_font(run, size: float = 10, bold: bool = False, color: str | None = None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def configure_docx(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Inches(1))
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color, before, after in [
        ("Title", 14, INK, 0, 8),
        ("Heading 1", 12, BLUE, 12, 5),
        ("Heading 2", 11, "2E74B5", 9, 4),
        ("Heading 3", 10, BLUE, 6, 3),
    ]:
        st = styles[name]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
    code = styles.add_style("CodeBlock", 1)
    code.font.name = FONT
    code._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(6)
    for sname in ["List Bullet", "List Number"]:
        st = styles[sname]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.font.size = Pt(10)
        st.paragraph_format.space_after = Pt(4)
    header = sec.header.paragraphs[0]
    header.text = "BIL302 Proje 3 - PicoRV32 RV32I FPGA Loader"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.runs[0], 8)
    footer = sec.footer.paragraphs[0]
    footer.text = "Courier New 10 pt akademik rapor"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.runs[0], 8)


def para(doc: Document, text: str = "", style: str | None = None):
    p = doc.add_paragraph(style=style)
    if text:
        set_run_font(p.add_run(text), 10)
    return p


def heading(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="CodeBlock")
    for line in text.strip("\n").splitlines():
        set_run_font(p.add_run(line + "\n"), 8.5)


def report_code_text(text: str, width: int = 96) -> str:
    out: list[str] = []
    for raw in text.rstrip().splitlines():
        expanded = raw.replace("\t", "    ")
        if len(expanded) <= width:
            out.append(expanded)
            continue
        indent = len(expanded) - len(expanded.lstrip(" "))
        prefix = " " * min(indent + 2, 12)
        parts = wrap(expanded, width=width, subsequent_indent=prefix, replace_whitespace=False, drop_whitespace=False)
        out.extend(parts or [""])
    return "\n".join(out)


def add_code_long(doc: Document, text: str, lines_per_block: int = 32) -> None:
    lines = report_code_text(text).splitlines()
    if not lines:
        add_code(doc, "")
        return
    for i in range(0, len(lines), lines_per_block):
        add_code(doc, "\n".join(lines[i:i + lines_per_block]))


def add_code_file(doc: Document, rel_path: str) -> None:
    path = ROOT / rel_path
    heading(doc, rel_path.replace("\\", "/"), 3)
    add_code_long(doc, path.read_text(encoding="utf-8"))


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, FILL)
        set_run_font(cell.paragraphs[0].add_run(h), 8.6, True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_run_font(cells[i].paragraphs[0].add_run(val), 8.4)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def add_figure(doc: Document, slug: str, caption: str) -> None:
    img = PNG_DIR / f"{slug}.png"
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(img), width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, 8.5)
    run.italic = True


def figure_title(slug: str) -> str:
    for fig in FIGURES:
        if fig.slug == slug:
            return fig.title
    raise KeyError(slug)


def build_docx() -> None:
    doc = Document()
    configure_docx(doc)
    t = doc.add_paragraph(style="Title")
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(t.add_run("PicoRV İşlemci Alt Kümesi (RV32I) için FPGA Tabanlı Loader Tasarımı"), 14, True, INK)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("BIL302 - 3. Proje/Tasarım Raporu\nTeslim: 07.06.2026\nGrup üyeleri: A.XX, B.YY, C.ZZ, D.WW"), 10, True)

    para(doc, "Bu rapor, önceki projelerde geliştirilen RV32I assembler ve linker modüllerini Tang Nano 9K üzerinde çalışan PicoRV32 tabanlı bir software-loader ile birleştiren uçtan uca sistemi açıklar. Çalışma; literatür araştırması, UART paket protokolü, checksum doğrulaması, host uygulaması, loader bellek haritası, deney tasarımı, sürdürülebilirlik etkileri, takım yönetimi ve bireysel katkı beyanlarını PDF yönergesindeki Program Çıktıları ile ilişkilendirir (PÇ1, PÇ6, PÇ7, PÇ8, PÇ12, PÇ13).")

    heading(doc, "1. Giriş ve Literatür Araştırması", 1)
    heading(doc, "1.1. Gömülü Sistemlerde Program Yükleme Mimarileri", 2)
    para(doc, "Gömülü sistemlerde bootloader yapıları; UART/SPI üzerinden program alma, JTAG ile hata ayıklama/programlama veya harici flash bellekteki imajı açılışta RAM'e taşıma gibi modellerle uygulanır. RISC-V'in açık ISA yapısı ve PicoRV32 gibi küçük soft-core çekirdekler, bu modelleri eğitim amaçlı FPGA sistemlerinde gözlenebilir kılar. UART loader düşük pin sayısı, standart PC desteği ve kolay paketlenebilir akış nedeniyle bu projede temel yöntem olarak seçilmiştir [1], [3], [4] (PÇ6).")
    para(doc, "JTAG güçlü hata ayıklama sağlar ancak sunum ve laboratuvar ortamında ek araç bağımlılığı yaratır. SPI flash kalıcı depolama sağlar fakat her denemede flash programlama süresi ve kart entegrasyonu gerekir. UART tabanlı software-loader ise FPGA bitstream'i sabit tutup yalnızca kullanıcı programını seri hattan yenileyerek geliştirme döngüsünü hızlandırır (PÇ6, PÇ8).")
    heading(doc, "1.2. Seri Haberleşme ve Veri Doğrulama Protokolleri", 2)
    para(doc, "UART asenkron bir seri protokoldür; 115200 baud 8N1 biçiminde her byte yaklaşık 10 bitlik hat süresi kullanır. Seri hatta byte kaybı veya bit bozulması olabileceği için host ile loader arasında veri doğrulama kodu gerekir. Checksum çok küçük kodla uygulanır; CRC ise burst hatalara karşı daha güçlüdür ancak loader kodunu büyütür [5], [6] (PÇ6, PÇ7).")
    add_table(doc, ["Yöntem", "Avantaj", "Dezavantaj", "Projede karar"], [
        ["8-bit checksum", "Çok küçük RV32I kodu, hızlı hesap", "Bazı hata örüntülerini kaçırabilir", "1 KiB boot alanına sığdığı için seçildi"],
        ["CRC-8/CRC-16", "Burst hatalara karşı daha güçlü", "Daha fazla kod ve çevrim maliyeti", "Gelecek sürüm alternatifi"],
        ["ACK/NACK", "Kayıpta yeniden gönderim altyapısı", "Protokol durumlarını artırır", "K/E/S cevaplarıyla temel geri bildirim sağlandı"],
    ], [1.3, 1.7, 1.7, 1.8])

    heading(doc, "2. Sistem Mimarisi ve Donanım-Yazılım Ortak Tasarımı", 1)
    heading(doc, "2.1. Toolchain Arayüz Standartları", 2)
    para(doc, "Sistem üç arayüz standardına ayrılmıştır: assembler-linker arasında RVOB1 relocatable object formatı, linker-host arasında .bin/.hex/.mem imajları, host-loader arasında ise adresli UART paket protokolü. Bu ayrım çok modüllü çalışmada ortak sözleşme oluşturur ve entegrasyon riskini azaltır (PÇ6, PÇ13).")
    add_figure(doc, "toolchain", figure_title("toolchain"))
    add_table(doc, ["Aşama", "Girdi", "Çıktı", "Sorumluluk"], [
        ["rvasm", "RV32I .s", "RVOB1 .ro", "Lexer, parser, encoding, relocation kaydı"],
        ["rvld", ".ro + link.toml", ".bin/.hex/.mem", "Section merge, sembol çözümleme, final image"],
        ["host_loader.py", ".bin veya .mem/.hex", "UART paketleri", "Adresli yazma, checksum ve START"],
        ["loader.s", "UART byte akışı", "BRAM yazma + jr entry", "PicoRV32 üzerinde yazılım tabanlı FSM"],
    ], [1.3, 1.6, 1.5, 2.1])
    heading(doc, "2.1.1. Compiler/Assembler Dizaynı ve Mimarisi", 3)
    para(doc, "Projede compiler rolünü eğitim amaçlı geliştirilen rvasm assembler aracı üstlenir. Klasik yüksek seviyeli compiler gibi C veya Java kaynak kodu işlemez; ancak kaynak metni aşamalı olarak token, AST, symbol table, machine code ve object format seviyelerine indirgediği için aynı derleyici mimarisi prensiplerini kullanır. Bu bölümde rvasm, RV32I assembly kaynak dosyasını donanımda çalışabilecek makine kodu parçalarına dönüştüren ön uç ve arka uç bileşenleriyle açıklanır (PÇ1, PÇ6).")
    add_figure(doc, "compiler_pipeline", figure_title("compiler_pipeline"))
    para(doc, "Compiler/assembler ön ucu lexer ve parser bileşenlerinden oluşur. Lexer, satır tabanlı assembly metnini IDENT, NUMBER, STRING, DIRECTIVE, REGISTER, COMMA ve COLON gibi tokenlara ayırır. Parser bu token akışını komut, etiket ve direktif düğümlerinden oluşan AST yapısına dönüştürür. Bu ayrım, hatalı sözdizimini erken yakalamayı ve sonraki encoding aşamasını sade tutmayı sağlar. Örneğin `main:` bir label düğümü, `li t0, 1` ise pseudo-instruction düğümü olarak modellenir (PÇ1).")
    para(doc, "Arka uç iki geçişli encoder olarak tasarlanmıştır. İlk geçişte .text ve .data section offsetleri hesaplanır, local/global/extern semboller symbol table'a yazılır ve her label için section-relative adres üretilir. İkinci geçişte gerçek RV32I instruction wordleri oluşturulur. R, I, S, B, U ve J tipindeki komutlarda opcode, rd, rs1, rs2, funct3, funct7 ve immediate alanları ilgili bit pozisyonlarına yerleştirilir. Böylece bit düzeyindeki ISA bilgisi tek bir encoder katmanında toplanır (PÇ1, PÇ6).")
    add_figure(doc, "compiler_data_model", figure_title("compiler_data_model"))
    add_table(doc, ["Compiler bileşeni", "Görevi", "Ürettiği veri"], [
        ["Lexer", "Kaynak metni token akışına çevirir", "Token listesi ve satır/kolon bilgisi"],
        ["Parser", "Tokenları AST düğümlerine dönüştürür", "Label, directive ve instruction AST'si"],
        ["Pass 1", "Section offsetlerini ve sembolleri hesaplar", "Symbol table, section boyutları"],
        ["Pass 2", "RV32I encoding ve pseudo expansion yapar", "Makine kodu byte dizisi"],
        ["RVOB1 writer", "Object dosyası header/table/payload alanlarını yazar", "Relocatable .ro dosyası"],
    ], [1.45, 2.6, 2.2])
    para(doc, "Relocation üretimi compiler/assembler mimarisinin en kritik parçasıdır. Kaynak dosya tek başına derlenirken extern sembollerin veya farklı object dosyalarındaki fonksiyonların final adresi bilinmez. Bu nedenle rvasm, doğrudan adres yazmak yerine relocation kaydı üretir. Örneğin `call delay` için PC-relative HI20/LO12 veya JAL tabanlı bir kayıt, `beq target` için branch relocation kaydı ve `.word symbol` için 32-bit absolute relocation kaydı tutulur. Bu karar, çok dosyalı programların bağımsız derlenmesini ve linker tarafından sonradan güvenli biçimde birleştirilmesini sağlar (PÇ1, PÇ7).")
    add_code(doc, """Compiler/assembler veri akışı:
1. Kaynak: examples/blink/blink.s
2. Lexer: IDENT(main), COLON, IDENT(li), REGISTER(t1), NUMBER(0x40000000)
3. Parser: Label(main), Instr(li), Instr(sw), Instr(call)
4. Pass 1: main=0x00, loop=0x18, extern delay
5. Pass 2: instruction bytes + R_RV32_PCREL_HI20/LO12 relocation
6. Çıktı: build/blink/blink.ro""")
    para(doc, "Hata yönetimi de compiler tasarımının parçasıdır. Bilinmeyen mnemonic, yanlış register adı, immediate aralık taşması, tekrar eden label veya eksik operand gibi hatalar assembly aşamasında durdurulur. Hatalı kaynak dosyanın object dosyasına dönüşmemesi, daha sonra linker veya FPGA aşamasında izlenmesi zor hataların oluşmasını engeller. Bu yaklaşım, sistem seviyesinde güvenilirlik için erken doğrulama ilkesini uygular (PÇ1, PÇ7).")
    heading(doc, "2.1.2. Linker Dizaynı ve Mimarisi", 3)
    para(doc, "rvld linker aracı, compiler/assembler tarafından üretilen bir veya daha fazla RVOB1 object dosyasını hedef bellek haritasına göre final çalıştırılabilir imaja dönüştürür. Linker tasarımının temel görevi, ayrı derlenen modüller arasında ortak adres uzayı kurmak, sembol referanslarını çözmek, relocation kayıtlarını uygulamak ve FPGA/host akışının beklediği .bin, .hex ve .mem formatlarını üretmektir (PÇ1, PÇ6).")
    add_figure(doc, "linker_pipeline", figure_title("linker_pipeline"))
    para(doc, "Linker mimarisi önce linker script dosyasını okuyarak ROM/RAM taban adreslerini ve section yerleşimini belirler. Ardından her object dosyasındaki .text sectionları giriş sırasına göre birleştirilir; .data sectionları ise seçilen layout politikasına göre final imajda uygun konuma eklenir. Her input section için linkOffset tutulur. Bu bilgi, object içindeki section-relative sembol değerlerini final mutlak adrese dönüştürmek için kullanılır (PÇ1).")
    para(doc, "Symbol resolution aşamasında local semboller yalnızca kendi object dosyası içinde görünür kalır; global semboller ortak tabloya eklenir; extern semboller ise bu global tabloya karşı çözümlenir. Aynı global sembolün iki dosyada tanımlanması duplicate symbol hatasıdır. Extern sembol bulunamazsa linker final imaj üretmez. Bu tasarım, eksik fonksiyon veya yanlış dosya sırası gibi sorunların FPGA'ya bozuk kod olarak yüklenmesini engeller (PÇ1, PÇ7).")
    add_table(doc, ["Linker aşaması", "Ana karar", "Hata kontrolü"], [
        ["Script okuma", "ROM/RAM taban adresleri ve entry noktası belirlenir", "Geçersiz adres veya eksik layout reddedilir"],
        ["Section merge", ".text ve .data final imaja taşınır", "Taşma ve hizalama kontrolleri yapılır"],
        ["Symbol table", "Global/local/extern semboller ayrıştırılır", "Duplicate ve undefined symbol hatası üretilir"],
        ["Relocation pass", "Instruction/data alanları final adrese göre patch edilir", "Range overflow hard error olur"],
        ["Image emit", ".bin/.hex/.mem dosyaları üretilir", "BRAM için 32-bit word formatı korunur"],
    ], [1.45, 2.75, 2.05])
    add_figure(doc, "relocation_flow", figure_title("relocation_flow"))
    para(doc, "Relocation pass, linker'ın instruction word üzerinde bit düzeyinde değişiklik yaptığı aşamadır. Her relocation kaydı için S sembol adresi, A addend ve P patch konumu hesaplanır. PC-relative türlerde genel ifade S + A - P biçimindedir. Branch relocation B-type immediate alanlarına, JAL relocation J-type immediate alanlarına, HI20/LO12 relocation ise lui/auipc ve addi/lw/sw gibi eşleşen komutların immediate bitlerine yazılır. Patch işlemi opcode ve register alanlarını korur; yalnızca hedef adresi temsil eden immediate alanlarını günceller (PÇ1, PÇ6).")
    para(doc, "Linker output katmanı üç farklı tüketiciye hizmet eder. `.bin` dosyası host loader tarafından seri porttan gönderilecek ham byte dizisidir. `.hex` dosyası Intel HEX benzeri metinsel kontrol formatıdır. `.mem` dosyası ise FPGA BRAM'in `$readmemh` çağrısıyla okuyabileceği 8 hex digitlik word satırlarından oluşur. Aynı final imajdan birden fazla çıktı üretmek, yazılım testleri ile FPGA demonstrasyonunun aynı binary gerçekliğe dayanmasını sağlar (PÇ7, PÇ13).")
    add_code(doc, """Linker örnek akışı:
rvld -script examples/blink/link.toml -o build/blink/blink start.ro blink.ro

1. _start, main ve delay sembolleri final adrese yerleştirilir.
2. call delay ve branch loop relocation kayıtları patch edilir.
3. build/blink/blink.bin, blink.hex ve blink.mem üretilir.
4. Eksik delay sembolü varsa final image üretilmez.""")
    para(doc, "Bu compiler-linker ayrımı, projenin loader hedefi için doğrudan önemlidir. Loader yalnızca final makine kodu wordlerini alır; ancak bu wordlerin doğru olması assembler'ın instruction encoding doğruluğuna ve linker'ın relocation doğruluğuna bağlıdır. Bu yüzden raporda compiler ve linker mimarisi ayrı ayrı açıklanmış, her iki katmanın çıktısı FPGA üzerinde çalışan loader zincirinin zorunlu ön koşulu olarak ele alınmıştır (PÇ1, PÇ6, PÇ7).")
    heading(doc, "2.2. FPGA Loader ve PicoRV32 Bellek Haritası", 2)
    para(doc, "FPGA açılışında PicoRV32 reset vektörü 0x00000000 adresindeki loader kodunu çalıştırır. Loader UART RX_READY yazmacını polling ile izler, paketleri alır, checksum doğruysa word hizalı olarak BRAM'deki kullanıcı program alanına yazar ve START paketinden sonra entry adresine jr komutuyla dallanır. Yönergede istenen FSM davranışı burada donanım FSM'i yerine PicoRV32 üzerinde çalışan yazılım FSM'i olarak gerçekleştirilmiştir (PÇ1, PÇ6, PÇ12).")
    add_figure(doc, "loader_fsm", figure_title("loader_fsm"))
    add_figure(doc, "memory_map", figure_title("memory_map"))
    add_table(doc, ["Adres", "İşlev", "Loader kullanımı"], [
        ["0x00000000-0x000003FF", "Boot loader alanı", "444 byte loader reset sonrası burada çalışır"],
        ["0x00000400-0x00001FFF", "Kullanıcı program alanı", "Host tarafından gönderilen word'ler buraya yazılır"],
        ["0x80000008 / 0x8000000C", "UART RX_DATA / RX_READY", "Gelen byte ve hazır bilgisi okunur"],
        ["0x80000010 / 0x80000014", "UART TX_DATA / TX_BUSY", "R, K, E, S cevapları gönderilir"],
    ], [2.1, 2.0, 2.2])
    add_figure(doc, "protocol", figure_title("protocol"))
    add_code(doc, "frame = 0x55, cmd, addr[0..3], words, data..., checksum\nchecksum = low8(cmd + addr bytes + words + data bytes)\ncmd=1 WRITE: word dizisi hedef adrese yazılır\ncmd=2 START: entry adresine dallanılır\nreplies: 'R'=ready, 'K'=write ok, 'E'=error, 'S'=starting")

    heading(doc, "3. Deneysel Çalışmalar, Test ve Analiz", 1)
    heading(doc, "3.1. Deney Tasarımı ve Test Senaryoları", 2)
    para(doc, "Test metodolojisi yalnızca derleme başarısını değil, FPGA üzerinde gözlenebilir I/O davranışını, linker sembol çözümlemesini ve loader hata kontrolünü ölçer. En az üç farklı Assembly programı seçilmiş; döngü, alt program çağrısı, .data kullanımı, UART çıktı ve hata senaryoları ayrı ayrı kapsanmıştır (PÇ7, PÇ12).")
    add_figure(doc, "test_coverage", figure_title("test_coverage"))
    add_table(doc, ["Senaryo", "Karmaşıklık", "Donanım gözlemi", "Kapsanan özellik"], [
        ["Blink / Knight Rider", "Döngü + delay çağrısı", "LED deseninin kayması", "branch, call, MMIO store"],
        ["Counter", "Alt program + sayaç", "LED'lerde sayaç değeri", "extern sembol, read_ms, set_leds"],
        ["UART Hello", ".data + string + UART", "Seri terminalde mesaj", "la pseudo, uart_puts"],
        ["Checksum negatif", "Hatalı paket", "Host 'E' cevabı alır", "veri doğrulama ve hata yönetimi"],
    ], [1.5, 1.5, 1.6, 1.9])
    add_code(doc, """.global main
.extern delay
main:
    li   t1, 0x40000000
    li   t2, 0x01
loop:
    sw   t2, 0(t1)
    li   a0, 0x40000
    call delay
    slli t2, t2, 1
    bne  t2, t3, loop""")
    add_code(doc, """counter_loop:
    call read_ms
    andi a0, a0, 0x3F
    call set_leds
    j    counter_loop""")
    add_code(doc, """uart_loop:
    la   a0, msg
    call uart_puts
    li   a0, 1000
    call delay_ms
    j    uart_loop
.data
msg: .asciz "Hello, RV32I!\\r\\n" """)
    heading(doc, "3.2. Veri Toplama ve Donanım Metrikleri", 2)
    para(doc, "Repository çıktılarında loader boyutu 444 byte, yani 111 adet 32-bit word olarak ölçülmüştür. Varsayılan host paketi 16 word/64 byte veri taşır. 115200 baud 8N1 için 64 byte veri içeren WRITE paketi yaklaşık 72 byte seri veri üretir; bu da paket başına yaklaşık 6,25 ms teorik hat süresi anlamına gelir. ACK ve USB-seri gecikmeleri gerçek süreyi artırır (PÇ7).")
    add_table(doc, ["Program boyutu", "Paket sayısı", "Teorik hat süresi", "Yorum"], [
        ["120 byte blink", "2", "~12,5 ms + ACK", "Küçük LED demosu"],
        ["260 byte UART hello", "5", "~31,3 ms + ACK", "String ve UART fonksiyonları"],
        ["512 byte", "8", "~50,0 ms + ACK", "Orta ölçekli kullanıcı programı"],
        ["4096 byte", "64", "~400 ms + ACK", "Kullanıcı alanının büyük bölümü"],
    ], [1.4, 1.2, 1.8, 2.1])
    add_table(doc, ["Kaynak metriği", "Ölçüm yeri", "Durum", "Not"], [
        ["LUT", "Gowin synthesis utilization report", "Sentez raporundan doldurulacak", "PicoRV32 + UART + decode mantığı"],
        ["Register", "Gowin synthesis report", "Sentez raporundan doldurulacak", "CPU, UART ve kontrol kayıtları"],
        ["BRAM/BSRAM", "Gowin synthesis report", "8 KiB program/data belleği", "bram_dp.v içinde $readmemh"],
        ["Fmax", "Timing report", "27 MHz üstü yeterli", "Tang Nano 9K sistem saati 27 MHz"],
    ], [1.5, 2.1, 1.5, 1.7])
    para(doc, "Gowin sentez raporu repository içinde bulunmadığı için LUT/Register/BRAM yüzdeleri tahmin edilmemiştir. Rapor, ölçülmüş yazılım metriklerini tahmini donanım metrikleriyle karıştırmamak için bu alanı açık ölçüm planı olarak bırakır; sunum öncesinde sentez çıktısı tabloya işlenmelidir (PÇ7, PÇ13).")

    heading(doc, "4. Projenin Küresel, Toplumsal ve Ekonomik Etkileri", 1)
    heading(doc, "4.1. Sürdürülebilirlik ve Yeşil Bilişim", 2)
    para(doc, "Software-loader yaklaşımı her kullanıcı programı değişikliğinde FPGA bitstream üretme zorunluluğunu azaltır. Daha kısa geliştirme döngüsü bilgisayar ve FPGA programlama süresini düşürür; küçük PicoRV32 çekirdeği ve sade UART çevre birimi de eğitim laboratuvarında gereksiz kaynak kullanımını sınırlar. Bu etki SKA 7 temiz enerji ve SKA 13 iklim eylemi ile ilişkilidir (PÇ8).")
    heading(doc, "4.2. Ekonomik Sürdürülebilirlik ve Teknolojik Bağımsızlık", 2)
    para(doc, "Açık RISC-V ekosistemi ve özgün assembler-linker-loader zinciri, kapalı kaynak mimarilere lisans ve araç bağımlılığını azaltır. Öğrenciler hazır binary araçları yalnızca kullanmak yerine instruction encoding, object format ve relocation mantığını tasarlayarak yerli gömülü sistem Ar-Ge'si için kritik bilgi birikimi kazanır. Bu yön SKA 8 ve SKA 9 hedeflerini destekler (PÇ6, PÇ8).")
    heading(doc, "4.3. Fonksiyonel Güvenlik ve Sağlık", 2)
    para(doc, "Tıbbi cihaz, otomotiv ve savunma gibi kritik sistemlerde hatalı firmware yükleme doğrudan güvenlik riski doğurur. Bu projedeki checksum ve START doğrulaması eğitim ölçeğinde basit olsa da temel ilkeyi gösterir: doğrulanmamış veri belleğe yazılsa bile çalıştırılmamalıdır. Daha kritik uygulamalarda CRC, imzalı firmware ve güvenli boot zinciri eklenmelidir (PÇ1, PÇ8).")
    heading(doc, "4.4. E-Atık Yönetimi ve Döngüsel Ekonomi", 2)
    para(doc, "Loader üzerinden seri portla program yenilemek, donanımı değiştirmeden yeni işlevlerin denenmesini sağlar. Bu yaklaşım eğitim kartlarının ömrünü uzatır, aynı kartın farklı deneylerde tekrar kullanılmasını kolaylaştırır ve elektronik atık oluşumunu azaltır. Bu etki SKA 12 sorumlu tüketim ve üretim hedefiyle uyumludur (PÇ8).")

    heading(doc, "5. Proje Yönetimi ve Takım Çalışması", 1)
    heading(doc, "5.1. Görev Dağılımı ve Sorumluluk Matrisi", 2)
    add_table(doc, ["İş paketi", "Responsible", "Accountable", "Consulted", "Informed"], [
        ["RV32I loader.s", "A.XX", "A.XX", "B.YY", "C.ZZ, D.WW"],
        ["Python host loader", "B.YY", "B.YY", "A.XX", "C.ZZ, D.WW"],
        ["UART/MMIO RTL entegrasyonu", "C.ZZ", "C.ZZ", "A.XX", "B.YY, D.WW"],
        ["Assembler/linker akışı", "D.WW", "D.WW", "A.XX, B.YY", "C.ZZ"],
        ["Test ve video", "Tüm ekip", "C.ZZ", "Tüm ekip", "Tüm ekip"],
        ["Raporlama", "Tüm ekip", "A.XX", "B.YY, C.ZZ, D.WW", "Tüm ekip"],
    ], [1.6, 1.15, 1.15, 1.15, 1.25])
    heading(doc, "5.2. Koordinasyon ve Sürüm Kontrol Yönetimi", 2)
    para(doc, "Kod, rapor, test ve FPGA dosyaları aynı Git repository içinde tutulmuştur. Ortak arayüz belgeleri; bellek haritası, UART paket formatı ve linker çıktı dosyalarıdır. Entegrasyon sonrası make loader, make demos ve go test ./... çalıştırılarak farklı modüllerin aynı protokol üzerinde uyumlu kaldığı doğrulanmıştır (PÇ12, PÇ13).")

    heading(doc, "6. Bireysel Katkı Beyanı", 1)
    para(doc, "Bu bölüm her öğrenci tarafından kendi gerçek katkısına göre doldurulup imzalanmalıdır. Her beyan, 'hangi spesifik modülü tamamen kendi başıma tasarladım ve tek başıma çözdüğüm en büyük teknik problem neydi?' sorusuna cevap vermelidir (PÇ12).")
    add_table(doc, ["Öğrenci", "Bağımsız tasarlanan modül", "Tek başına çözülen problem", "İmza"], [
        ["A.XX", "RV32I software-loader ve loader_link.toml", "Checksum doğrulayıp BRAM'e word hizalı yazma", ""],
        ["B.YY", "host_loader.py paketleme akışı", ".bin/.mem dosyasını adresli UART paketlerine dönüştürme", ""],
        ["C.ZZ", "UART MMIO ve top entegrasyonu", "RX_READY/TX_BUSY kayıtlarını PicoRV32 arayüzüne bağlama", ""],
        ["D.WW", "Assembler/linker entegrasyonu ve rapor doğrulaması", "RVOB1 çıktısı ile final .mem imajı arasındaki sembol/relocation akışını doğrulama", ""],
    ], [1.1, 2.0, 2.6, 0.8])

    heading(doc, "7. Kaynakça", 1)
    for ref in [
        '[1] RISC-V International, "The RISC-V Instruction Set Manual, Volume I: Unprivileged ISA."',
        '[2] J. R. Levine, Linkers and Loaders, Morgan Kaufmann, 1999.',
        '[3] YosysHQ, "PicoRV32 - A Size-Optimized RISC-V CPU."',
        '[4] RISC-V International, "RISC-V ELF psABI Specification."',
        '[5] P. Koopman, "32-Bit Cyclic Redundancy Codes for Internet Applications," DSN, 2002.',
        '[6] Nandland, "UART, Serial Port, RS-232 Interface" UART implementation notes.',
        '[7] United Nations, "Sustainable Development Goals."',
        '[8] Gowin Semiconductor, "Tang Nano 9K / GW1NR FPGA Documentation."',
        '[9] GNU Binutils Documentation, "Assembler and Linker Documentation."',
        '[10] Proje kaynak kodları: assembler/, linker/, gowin/loader.s, scripts/host_loader.py, docs/*.md.',
    ]:
        para(doc, ref)

    heading(doc, "Ek A - Değerlendirme Kriterleri Karşılama Matrisi", 1)
    add_table(doc, ["Kriter", "Rapordaki kanıt", "PÇ"], [
        ["Literatür araştırması", "UART/SPI/JTAG ve checksum/CRC karşılaştırması", "PÇ6"],
        ["Yöntem ve tasarım", "Toolchain arayüzleri, bellek haritası, FSM ve alternatif gerekçesi", "PÇ6"],
        ["Test tasarımı", "Blink, counter, UART hello ve checksum negatif testi", "PÇ7"],
        ["Analiz ve yorum", "Loader boyutu, paket süresi ve kaynak metrik planı", "PÇ7"],
        ["Sürdürülebilirlik", "Enerji, Ar-Ge maliyeti, güvenlik ve e-atık etkileri", "PÇ8"],
        ["Takım çalışması", "RACI ve sürüm kontrol yönetimi", "PÇ12/PÇ13"],
        ["Loader doğruluğu", "Checksum, ACK/ERR, START ve BRAM yazma akışı", "PÇ1"],
        ["FPGA örnek uygulama", "LED ve UART gözlenebilir senaryolar", "PÇ12"],
        ["Rapor kalitesi", "Courier New 10 pt, SVG kaynaklı şekiller, tablo ve kod blokları", "PÇ13"],
    ], [1.6, 3.8, 1.0])

    heading(doc, "Ek B - Teslim ve Video Notu", 1)
    para(doc, "Teslim PDF dosya adı yönergeye uygun olarak BIL302_PROJE3_A.XX_B.YY_C.ZZ_D.WW_170526.PDF biçiminde üretilmiştir. Kısa video için sistem çalışmasının yalnızca çıktı göstermemesi; host script, paket protokolü, loader akışı ve FPGA üzerindeki LED/UART gözlemini en fazla 5 dakikada teknik sunum kalitesinde anlatması önerilir (PÇ12, PÇ13).")
    add_code(doc, "make loader\nmake demos\ngo test ./...\npython scripts\\host_loader.py build\\blink\\blink.bin -p COM5 --base 0x400 --entry 0x400")
    heading(doc, "Ek C - Tüm Example Programları ve Linker Scriptleri", 1)
    para(doc, "Bu ekte repository içindeki tüm example dosyaları tam kaynak kodlarıyla verilmiştir. Blink örneği LED MMIO ve delay çağrısını, counter örneği alt program ve sayaç tabanlı I/O akışını, UART örneği ise .data section, string kullanımı ve seri çıkış fonksiyonlarını kapsar. Bu kodlar, 3.1 bölümünde açıklanan deney senaryolarının doğrudan uygulanabilir kaynaklarıdır (PÇ7, PÇ12).")
    example_files = [
        "examples/blink/start.s",
        "examples/blink/blink.s",
        "examples/blink/link.toml",
        "examples/counter/start.s",
        "examples/counter/io.s",
        "examples/counter/counter.s",
        "examples/counter/link.toml",
        "examples/uart/start.s",
        "examples/uart/io.s",
        "examples/uart/uart.s",
        "examples/uart/hello.s",
        "examples/uart/link.toml",
    ]
    add_table(doc, ["Example dosyası", "Rolü"], [
        ["examples/blink/start.s", "Blink programı için reset/start kodu ve delay fonksiyonu"],
        ["examples/blink/blink.s", "LED desenini döngü içinde değiştiren ana blink uygulaması"],
        ["examples/blink/link.toml", "Blink final imajının bellek yerleşimi"],
        ["examples/counter/start.s", "Counter örneği başlangıç kodu"],
        ["examples/counter/io.s", "Counter için read_ms ve set_leds I/O yardımcıları"],
        ["examples/counter/counter.s", "Sayaç değerini LED çıkışına bağlayan ana program"],
        ["examples/counter/link.toml", "Counter final imajının bellek yerleşimi"],
        ["examples/uart/start.s", "UART örneği başlangıç kodu"],
        ["examples/uart/io.s", "UART örneği zamanlama/I/O yardımcı kodu"],
        ["examples/uart/uart.s", "UART putc/puts fonksiyonları"],
        ["examples/uart/hello.s", "Hello mesajını seri porta yazan ana program"],
        ["examples/uart/link.toml", "UART final imajının bellek yerleşimi"],
    ], [2.4, 4.1])
    for rel in example_files:
        add_code_file(doc, rel)

    heading(doc, "Ek D - FPGA Loader Assembly Kodu", 1)
    para(doc, "Aşağıdaki `gowin/loader.s` dosyası, FPGA açılışında PicoRV32 üzerinde çalışan software-loader kodudur. UART paketlerini okur, checksum doğrular, WRITE paketlerinde BRAM'e word yazar ve START paketinde kullanıcı programına dallanır (PÇ1, PÇ7, PÇ12).")
    add_code_file(doc, "gowin/loader.s")

    heading(doc, "Ek E - Host Loader Python Kodu", 1)
    para(doc, "Aşağıdaki `scripts/host_loader.py` dosyası, linker çıktısı olan .bin/.mem/.hex program imajını okuyup seri porttan FPGA loader'a paketler halinde gönderen host uygulamasıdır. Paketleme, checksum üretimi, ACK/ERR kontrolü ve START komutu bu dosyada uygulanır (PÇ1, PÇ7, PÇ13).")
    add_code_file(doc, "scripts/host_loader.py")
    doc.save(DOCX)


def pdf_styles():
    pdfmetrics.registerFont(TTFont("CourierNew", r"C:\Windows\Fonts\cour.ttf"))
    pdfmetrics.registerFont(TTFont("CourierNew-Bold", r"C:\Windows\Fonts\courbd.ttf"))
    pdfmetrics.registerFont(TTFont("CourierNew-Italic", r"C:\Windows\Fonts\couri.ttf"))
    return {
        "body": ParagraphStyle("body", fontName="CourierNew", fontSize=10, leading=12, spaceAfter=6, alignment=TA_LEFT),
        "title": ParagraphStyle("title", fontName="CourierNew-Bold", fontSize=14, leading=17, spaceAfter=10, alignment=TA_CENTER, textColor=colors.HexColor("#0B2545")),
        "h1": ParagraphStyle("h1", fontName="CourierNew-Bold", fontSize=12, leading=15, spaceBefore=10, spaceAfter=5, textColor=colors.HexColor("#1F4D78")),
        "h2": ParagraphStyle("h2", fontName="CourierNew-Bold", fontSize=11, leading=14, spaceBefore=8, spaceAfter=4, textColor=colors.HexColor("#2E74B5")),
        "h3": ParagraphStyle("h3", fontName="CourierNew-Bold", fontSize=10, leading=12, spaceBefore=6, spaceAfter=3, textColor=colors.HexColor("#1F4D78")),
        "code": ParagraphStyle("code", fontName="CourierNew", fontSize=8.4, leading=10, leftIndent=12, spaceAfter=6),
        "cell": ParagraphStyle("cell", fontName="CourierNew", fontSize=7.3, leading=8.8, spaceAfter=0),
        "cell_bold": ParagraphStyle("cell_bold", fontName="CourierNew-Bold", fontSize=7.3, leading=8.8, spaceAfter=0),
        "caption": ParagraphStyle("caption", fontName="CourierNew-Italic", fontSize=8.5, leading=10, alignment=TA_CENTER, spaceAfter=8),
    }


def add_pdf_table(story, styles, headers, rows, widths=None):
    data = [[Paragraph(escape(c), styles["cell_bold"]) for c in headers]]
    data.extend([[Paragraph(escape(c), styles["cell"]) for c in row] for row in rows])
    total = 6.5 * inch
    if widths:
        s = sum(widths)
        col_widths = [total * w / s for w in widths]
    else:
        col_widths = [total / len(headers)] * len(headers)
    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign="CENTER")
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#999999")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(table)
    story.append(Spacer(1, 8))


def build_pdf_from_docx() -> None:
    styles = pdf_styles()
    doc = Document(DOCX)
    story = []
    fig_index = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = next((p for p in doc.paragraphs if p._p is child), None)
            if paragraph is None:
                continue
            text = paragraph.text.strip()
            if not text:
                continue
            name = paragraph.style.name if paragraph.style else ""
            if text.startswith("Şekil "):
                story.append(KeepTogether([PdfImage(str(PNG_DIR / f"{FIGURES[fig_index].slug}.png"), width=6.2 * inch, height=3.1 * inch), Paragraph(escape(text), styles["caption"])]))
                fig_index += 1
            elif name == "Title":
                story.append(Paragraph(escape(text).replace("\n", "<br/>"), styles["title"]))
            elif name == "Heading 1":
                story.append(Paragraph(escape(text), styles["h1"]))
            elif name == "Heading 2":
                story.append(Paragraph(escape(text), styles["h2"]))
            elif name == "Heading 3":
                story.append(Paragraph(escape(text), styles["h3"]))
            elif name == "CodeBlock":
                story.append(Preformatted(text, styles["code"], maxLineLength=96))
            else:
                story.append(Paragraph(escape(text).replace("\n", "<br/>"), styles["body"]))
        elif child.tag == qn("w:tbl"):
            table = next((t for t in doc.tables if t._tbl is child), None)
            if table is None:
                continue
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            add_pdf_table(story, styles, rows[0], rows[1:])
    def page(canvas, pdf_doc):
        canvas.saveState()
        canvas.setFont("CourierNew", 8)
        canvas.drawString(inch, 0.55 * inch, "BIL302 Proje 3 - RV32I FPGA Loader")
        canvas.drawRightString(7.5 * inch, 0.55 * inch, f"Sayfa {pdf_doc.page}")
        canvas.restoreState()
    pdf = SimpleDocTemplate(str(PDF), pagesize=letter, rightMargin=inch, leftMargin=inch, topMargin=inch, bottomMargin=inch)
    pdf.build(story, onFirstPage=page, onLaterPages=page)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    write_figures()
    build_docx()
    build_pdf_from_docx()
    print(DOCX)
    print(PDF)


if __name__ == "__main__":
    main()
